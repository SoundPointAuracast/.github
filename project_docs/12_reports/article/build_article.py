#!/usr/bin/env python3
"""Build the final scientific article (DOCX) for the Route+ / «Точка звука» project.

Fallback layout: A4, Times New Roman 12, single spacing, 20 mm margins,
1.25 cm first-line indent. Page limit: 1.5–2 pages (checked via PDF).
"""
from __future__ import annotations

import os

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(HERE, "final", "Статья_Точка_звука_Никитин_ДК_FINAL.docx")

TITLE = ("Разработка локальной ассистивной аудиосистемы с синхронизированным "
         "текстовым сопровождением для образовательных пространств")
AUTHOR = "Никитин Даниил Константинович"
AFFIL = "МГТУ им. Н.Э. Баумана, Москва"

ABSTRACT = (
    "В образовательных аудиториях расстояние, шум и реверберация снижают "
    "разборчивость речи, а для пользователей слуховых аппаратов и кохлеарных "
    "имплантов потеря части фразы означает потерю учебной информации. Существующие "
    "системы персональной передачи речи доставляют звук, но не формируют текстовый "
    "слой занятия. В работе рассматривается архитектура, где единый чистый речевой "
    "источник питает персональный аудиоканал и синхронизированное текстовое "
    "сопровождение; описаны программный прототип Route+ и программа "
    "экспериментальной проверки."
)

KEYWORDS = ("Auracast; Bluetooth LE Audio; инклюзивное образование; ассистивные "
            "технологии; кохлеарный имплант; распознавание речи; доступность")

SECTIONS: list[tuple[str, list[str]]] = [
    ("Актуальность и постановка задачи", [
        "В лекционной аудитории речевой сигнал приходит ослабленным, зашумлённым и "
        "реверберированным; для человека с нарушением слуха это не вопрос комфорта, а "
        "потеря смысла. Особенно остра проблема на дальних рядах и при большой "
        "реверберации: исследования показывают, что даже при нормативном уровне шума "
        "отношение сигнал/шум около +15 дБ может быть недостаточным для младших "
        "школьников [6]. Возникает вопрос: как обеспечить пользователю слухового "
        "аппарата (СА) или кохлеарного импланта (КИ) стабильный одновременный доступ "
        "к звуку и смыслу речи?",
    ]),
    ("Существующие решения", [
        "В российской образовательной практике применяются FM-радиоклассы, в "
        "частности «СОНЕТ 2.0» (ГК «Исток-Аудио»): микрофон преподавателя, передатчик "
        "и приёмники с индукционным выходом (индуктор или петля) для работы СА и КИ в "
        "режиме T/TM; поддерживается групповое использование [4].",
        "Проприетарные цифровые системы (Phonak Roger) используют удалённый микрофон "
        "и приёмники для СА и КИ в диапазоне 2,4 ГГц с заявленной задержкой передачи "
        "менее 20 мс. Существенно, что универсальный приёмник Roger NeckLoop совмещает "
        "индукционный выход (T-coil) с USB-аудиоинтерфейсом: официальная инструкция "
        "производителя описывает использование микрофона Roger как входа для "
        "сторонних систем распознавания речи и генерации субтитров (2021) [3]. Идея "
        "«чистый фид → распознавание речи» не является новой и новизной не заявляется.",
        "Открытый стандарт Auracast (Bluetooth LE Audio) обеспечивает "
        "широковещательную передачу аудио на совместимые устройства локально, без "
        "интернета; смартфон не входит в аудиотракт [1, 2]. Внедрения в вузах "
        "(University of Queensland — 65 аудиторий; Oxford; UAL) подтверждают "
        "работоспособность аудиотракта; текстовый слой занятия в них не подтверждён "
        "[7]. Принцип проекта: для широкой аудитории это удобство, для человека с "
        "нарушением слуха — доступ к информации.",
    ]),
    ("Предлагаемая архитектура", [
        "В работе рассматривается архитектура «Точка звука»: единый чистый речевой "
        "источник (микрофон преподавателя или выход микшера) разветвляется до "
        "кодирования LC3 на две ветки.",
        "Аудиоветка: PCM → LC3 → Auracast → совместимые персональные устройства. "
        "Базовая конфигурация LC3 — 24_2_1 (24 кГц, 10 мс, 48 кбит/с); приёмник "
        "обязан поддерживать 16/24 кГц и presentation delay 40 мс [1, 2]. Выбор "
        "24_2_1 — инженерный baseline, подлежащий проверке. Для устройств без "
        "Auracast предусмотрен Bridge-T: Auracast Sink → декодирование LC3 → PCM → "
        "ЦАП → усилитель → индукционная катушка → режим T/MT. Класс подобных "
        "устройств известен (Roger MyLink/NeckLoop, Auri RX1, Bettear RTX, AuraCoil), "
        "поэтому Bridge-T не заявляется как новизна.",
        "Текстовая ветка: PCM → ASR → транскрипт с временными метками → Route+. "
        "Приложение не входит в аудиотракт Auracast и не управляет широковещательным "
        "подключением — это ограничение мобильных операционных систем [5]. Обе ветки "
        "разделяют общий временной контекст сессии, что позволяет синхронизировать "
        "текст с аудио и восстанавливать пропущенные фрагменты.",
    ]),
    ("Программный MVP и методика проверки", [
        "Разработан программный прототип Route+ (React + FastAPI, WebSocket). "
        "Реализованы: вход в сессию, поток субтитров (partial/final), временные "
        "метки, функция «Не расслышал» — дословный текст последних 15 секунд "
        "(значение по умолчанию; оптимум не доказан), история, заметки, закладки, "
        "настройки доступности и демонстрационный режим без микрофона и "
        "Auracast-оборудования. В демонстрационной конфигурации распознавание "
        "выполняет заглушка (MockASRProvider); реальный локальный ASR (T-one, "
        "GigaAM-v3 как эталон) — следующий этап.",
        "Разработана программа проверки: E01 — сквозная задержка аудиотракта (цель "
        "≤60 мс, допустимо ≤100 мс; TARGET, измерения не выполнены); E03 — "
        "устойчивость передачи; E04 — сравнение конфигураций LC3; E05 — сравнение "
        "точности распознавания на чистом фиде и микрофоне смартфона (WER, CER, "
        "задержки субтитров и финализации); E07 — передача через Bridge-T. Ключевой "
        "программный эксперимент — E05.",
    ]),
    ("Исследуемое отличие", [
        "В работе исследуется системный подход, при котором единый чистый речевой "
        "источник используется одновременно для персонального аудиоканала и "
        "синхронизированного цифрового сопровождения образовательной сессии. Новизна "
        "отдельных элементов (clean-feed, Auracast, индукционный мост, субтитры) не "
        "заявляется. Предмет дальнейших исследований — общий таймлайн аудио и текста, "
        "привязка транскрипта к сессии, извлечение пропущенного фрагмента. Патентная "
        "новизна не утверждается до профессионального патентного поиска.",
    ]),
    ("Заключение", [
        "Разработаны архитектура и программный прототип Route+; определена программа "
        "экспериментальной проверки. Следующий этап — сборка аппаратного стенда "
        "(рекомендован Nordic nRF5340 Audio DK ×2), реализация локального ASR и "
        "выполнение экспериментов E01, E03, E04, E05, E07. Утверждения об "
        "эффективности будут сформулированы только по результатам измерений.",
    ]),
]

REFERENCES = [
    "Bluetooth SIG. Public Broadcast Profile 1.0.1 [Электронный ресурс]. "
    "URL: https://www.bluetooth.com/specifications/specs/public-broadcast-profile/ "
    "(дата обращения: 26.09.2026).",
    "Bluetooth SIG. Basic Audio Profile 1.0.1 [Электронный ресурс]. "
    "URL: https://www.bluetooth.com/specifications/specs/basic-audio-profile-1-0-1/ "
    "(дата обращения: 26.09.2026).",
    "Phonak. Roger NeckLoop: USB audio and speech-to-text [Электронный ресурс]. "
    "URL: https://www.phonak.com/en-us/hearing-devices/microphones/roger-neckloop "
    "(дата обращения: 26.09.2026).",
    "Радиокласс «СОНЕТ 2.0» [Электронный ресурс] // Исток-Аудио. "
    "URL: https://www.istok-audio.com/catalog/product/radioklass_sonet_2_0/ "
    "(дата обращения: 26.09.2026).",
    "Bluetooth LE Audio [Электронный ресурс] // Android Developers. "
    "URL: https://developer.android.com/develop/connectivity/bluetooth/ble-audio/overview "
    "(дата обращения: 26.09.2026).",
    "Bradley J. S., Sato H. The intelligibility of speech in elementary school "
    "classrooms // The Journal of the Acoustical Society of America. 2008. Vol. 123, "
    "№ 4. P. 2078–2086. DOI: 10.1121/1.2839285.",
    "UQ creates change for accessibility students [Электронный ресурс] // "
    "University of Queensland. 2026. "
    "URL: https://news.uq.edu.au/2026-03-uq-creates-change-accessibility-students "
    "(дата обращения: 26.09.2026).",
]


def set_normal_style(doc: Document) -> None:
    st = doc.styles["Normal"]
    st.font.name = "Times New Roman"
    st.font.size = Pt(12)
    rpr = st.element.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    rfonts.set(qn("w:eastAsia"), "Times New Roman")
    rfonts.set(qn("w:cs"), "Times New Roman")
    pf = st.paragraph_format
    pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    pf.line_spacing = 1.0
    pf.space_after = Pt(0)
    pf.space_before = Pt(0)
    pf.first_line_indent = Cm(1.25)


def para(doc, runs, size=12, align=WD_ALIGN_PARAGRAPH.JUSTIFY, bold=False,
         italic=False, indent=Cm(1.25), before=0, after=0, hanging=None):
    p = doc.add_paragraph()
    p.alignment = align
    pf = p.paragraph_format
    pf.first_line_indent = indent
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = 1.0
    if hanging is not None:
        pf.left_indent = hanging
        pf.first_line_indent = Cm(-hanging.cm)
    if isinstance(runs, str):
        runs = [(runs, bold, italic)]
    for text, b, i in runs:
        r = p.add_run(text)
        r.font.name = "Times New Roman"
        r.font.size = Pt(size)
        r.bold = b
        r.italic = i
        rpr = r._element.get_or_add_rPr()
        rf = rpr.get_or_add_rFonts()
        rf.set(qn("w:eastAsia"), "Times New Roman")
        rf.set(qn("w:cs"), "Times New Roman")
    return p


def build() -> Document:
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.top_margin = sec.bottom_margin = Cm(2)
    sec.left_margin = sec.right_margin = Cm(2)
    set_normal_style(doc)

    para(doc, TITLE, size=14, align=WD_ALIGN_PARAGRAPH.CENTER, bold=True,
         indent=Cm(0), after=3)
    para(doc, AUTHOR, size=12, align=WD_ALIGN_PARAGRAPH.CENTER, indent=Cm(0), after=1)
    para(doc, AFFIL, size=11, align=WD_ALIGN_PARAGRAPH.CENTER, italic=True,
         indent=Cm(0), after=6)

    para(doc, [("Аннотация. ", True, False), (ABSTRACT, False, False)],
         size=11, indent=Cm(1.25), after=3)
    para(doc, [("Ключевые слова: ", True, False), (KEYWORDS, False, False)],
         size=11, indent=Cm(1.25), after=6)

    for heading, paragraphs in SECTIONS:
        para(doc, f"{heading}", size=12, align=WD_ALIGN_PARAGRAPH.LEFT, bold=True,
             indent=Cm(0), before=2, after=1)
        for txt in paragraphs:
            para(doc, txt, size=12, after=0)

    para(doc, "Литература", size=12, align=WD_ALIGN_PARAGRAPH.LEFT, bold=True,
         indent=Cm(0), before=2, after=1)
    for i, ref in enumerate(REFERENCES, 1):
        para(doc, f"{i}. {ref}", size=11, indent=Cm(0), after=0, hanging=Cm(0.75))

    os.makedirs(os.path.dirname(OUT), exist_ok=True)
    doc.save(OUT)
    return doc


def plain_text() -> str:
    lines = [TITLE, "", AUTHOR, AFFIL, "", "Аннотация. " + ABSTRACT, "",
             "Ключевые слова: " + KEYWORDS, ""]
    for heading, paragraphs in SECTIONS:
        lines.append(heading)
        lines.extend(paragraphs)
        lines.append("")
    lines.append("Литература")
    lines.extend(f"{i}. {r}" for i, r in enumerate(REFERENCES, 1))
    return "\n".join(lines)


def counts() -> None:
    text = plain_text()
    words = len(text.split())
    print(f"words={words} chars_no_spaces={sum(1 for c in text if not c.isspace())} "
          f"chars_with_spaces={len(text)}")
    plain_path = os.path.join(HERE, "final", "Статья_Точка_звука_Никитин_ДК_PLAIN.txt")
    with open(plain_path, "w", encoding="utf-8") as f:
        f.write(text + "\n")
    print("plain:", plain_path)


if __name__ == "__main__":
    build()
    counts()
    print("saved:", OUT)
