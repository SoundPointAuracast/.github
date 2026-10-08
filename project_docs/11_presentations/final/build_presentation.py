#!/usr/bin/env python3
"""Build the final Route+ / «Точка звука» presentation (editable PPTX).

Generator kept in the project so the deck can be rebuilt.
Output: 11_presentations/final/Разработка_локального_аудиоретранслятора_FINAL.pptx
"""
from __future__ import annotations

import os

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.abspath(os.path.join(HERE, "..", ".."))
SHOTS = os.path.join(ROOT, "13_demo", "app_demo", "screenshots")
OUT = os.path.join(HERE, "Разработка_локального_аудиоретранслятора_FINAL.pptx")

# ---------------------------------------------------------------- palette
NAVY = RGBColor(0x12, 0x30, 0x5C)
BLUE = RGBColor(0x1F, 0x6F, 0xD0)
LIGHT = RGBColor(0xE7, 0xF0, 0xFB)
LIGHT2 = RGBColor(0xF4, 0xF6, 0xF9)
INK = RGBColor(0x16, 0x23, 0x3A)
MUTED = RGBColor(0x5B, 0x6B, 0x83)
LINEC = RGBColor(0xD7, 0xDE, 0xE8)
GREEN = RGBColor(0x1E, 0x7A, 0x46)
GREEN_L = RGBColor(0xE3, 0xF4, 0xE8)
ORANGE = RGBColor(0xB3, 0x5C, 0x00)
ORANGE_L = RGBColor(0xFD, 0xF0, 0xE0)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
BLUE_SOFT = RGBColor(0x9E, 0xC5, 0xF0)
FONT = "Arial"

SW, SH = 13.333, 7.5
M = 0.6


# ---------------------------------------------------------------- helpers
def new_deck() -> Presentation:
    prs = Presentation()
    prs.slide_width = Inches(SW)
    prs.slide_height = Inches(SH)
    return prs


def slide(prs: Presentation):
    return prs.slides.add_slide(prs.slide_layouts[6])


def rect(sl, x, y, w, h, fill=None, line=None, lw=1.0, radius=0.06,
         shape=MSO_SHAPE.ROUNDED_RECTANGLE):
    shp = sl.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    try:
        shp.shadow.inherit = False
    except Exception:
        pass
    if fill is None:
        shp.fill.background()
    else:
        shp.fill.solid()
        shp.fill.fore_color.rgb = fill
    if line is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = line
        shp.line.width = Pt(lw)
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        try:
            shp.adjustments[0] = radius
        except Exception:
            pass
    shp.text_frame.word_wrap = True
    return shp


def text(sl, x, y, w, h, paras, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    """paras: list of dicts: text | runs, size, bold, color, italic, after, before, line."""
    tb = sl.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    for i, p in enumerate(paras):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para.alignment = p.get("align", align)
        if "line" in p:
            para.line_spacing = p["line"]
        if "before" in p:
            para.space_before = Pt(p["before"])
        if "after" in p:
            para.space_after = Pt(p["after"])
        runs = p.get("runs") or [(p.get("text", ""), {})]
        for rt, ro in runs:
            r = para.add_run()
            r.text = rt
            f = r.font
            f.name = FONT
            f.size = Pt(ro.get("size", p.get("size", 18)))
            f.bold = ro.get("bold", p.get("bold", False))
            f.italic = ro.get("italic", p.get("italic", False))
            f.color.rgb = ro.get("color", p.get("color", INK))
    return tb


def box_text(sl, x, y, w, h, paras, fill, line=None, align=PP_ALIGN.CENTER,
             anchor=MSO_ANCHOR.MIDDLE, radius=0.08):
    rect(sl, x, y, w, h, fill=fill, line=line, radius=radius)
    text(sl, x + 0.08, y + 0.04, w - 0.16, h - 0.08, paras, align=align, anchor=anchor)


def arrow(sl, x1, y1, x2, y2, color=NAVY, width=1.5, dash=None, head=True):
    conn = sl.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1),
                                   Inches(x2), Inches(y2))
    conn.line.color.rgb = color
    conn.line.width = Pt(width)
    if dash is not None:
        conn.line.dash_style = dash
    if head:
        ln = conn.line._get_or_add_ln()
        tail = ln.makeelement(qn("a:tailEnd"), {"type": "triangle", "w": "med", "len": "med"})
        ln.append(tail)
    return conn


def badge(sl, x, y, w, h, label, fill, fg=WHITE, size=12):
    box_text(sl, x, y, w, h, [{"text": label, "size": size, "bold": True, "color": fg}],
             fill=fill, radius=0.5)


def header(sl, eyebrow, title, tsize=28):
    text(sl, M, 0.28, SW - 2 * M, 0.26,
         [{"text": eyebrow.upper(), "size": 12, "bold": True, "color": BLUE}])
    text(sl, M, 0.56, SW - 2 * M, 0.75,
         [{"text": title, "size": tsize, "bold": True, "color": NAVY, "line": 0.95}])


def footer(sl, txt):
    text(sl, M, 7.06, SW - 2 * M, 0.3, [{"text": txt, "size": 10, "color": MUTED}])


def picture_cover(sl, path, x, y, w, h):
    iw, ih = Image.open(path).size
    target = w / h
    source = iw / ih
    pic = sl.shapes.add_picture(path, Inches(x), Inches(y), width=Inches(w), height=Inches(h))
    if source > target:
        c = (1 - target / source) / 2
        pic.crop_left = c
        pic.crop_right = c
    elif source < target:
        c = (1 - source / target) / 2
        pic.crop_top = c
        pic.crop_bottom = c
    return pic


def card(sl, x, y, w, h, fill=WHITE, line=LINEC):
    rect(sl, x, y, w, h, fill=fill, line=line, lw=1.0, radius=0.07)


# ---------------------------------------------------------------- slides
def s01_title(prs):
    sl = slide(prs)
    rect(sl, 0, 0, SW, SH, fill=NAVY, shape=MSO_SHAPE.RECTANGLE)
    rect(sl, 0.92, 1.75, 0.075, 2.6, fill=BLUE, shape=MSO_SHAPE.RECTANGLE)

    text(sl, 1.25, 1.75, 10.8, 0.3,
         [{"text": "ИНКЛЮЗИВНЫЙ МАРШРУТ+ / ТОЧКА ЗВУКА", "size": 14, "bold": True,
           "color": BLUE_SOFT}])
    text(sl, 1.25, 2.15, 10.9, 1.75,
         [{"text": "Локальный аудиоретранслятор персональной речи", "size": 38,
           "bold": True, "color": WHITE, "line": 1.0}])
    text(sl, 1.25, 4.0, 10.9, 0.8,
         [{"text": "Единая образовательная сессия: персональный звук и синхронизированный текст",
           "size": 19, "color": LIGHT, "line": 1.1}])

    badge(sl, 1.25, 4.95, 3.4, 0.45, "Route+ MVP — реализован", GREEN, WHITE, 13)
    badge(sl, 4.85, 4.95, 4.6, 0.45, "Hardware-стенд — рекомендован, не закуплен",
          ORANGE, WHITE, 13)
    badge(sl, 9.65, 4.95, 3.0, 0.45, "Эксперименты — PLANNED", ORANGE, WHITE, 13)

    text(sl, 1.25, 6.35, 11.0, 0.35,
         [{"text": "Автор: Никитин Даниил Константинович · 2026 · GATE 6B",
           "size": 14, "color": BLUE_SOFT}])


def s02_motivation(prs):
    sl = slide(prs)
    header(sl, "Личная мотивация автора", "Откуда идея")
    card(sl, 0.9, 1.65, 11.53, 2.7, fill=LIGHT2)
    rect(sl, 0.9, 1.65, 0.08, 2.7, fill=BLUE, shape=MSO_SHAPE.RECTANGLE)
    text(sl, 1.25, 1.95, 10.9, 2.1,
         [{"text": "«Я — пользователь кохлеарного импланта. Во время обучения я "
                   "пользовался персональной радиосистемой преподавателя и знаю, "
                   "как много зависит от того, слышишь ли ты речь целиком».", "size": 24,
           "color": INK, "line": 1.15}])
    card(sl, 0.9, 4.7, 5.6, 1.7)
    text(sl, 1.15, 4.95, 5.1, 1.25,
         [{"text": "Идея — из личного опыта", "size": 18, "bold": True, "color": NAVY},
          {"text": "Не заменить существующие системы, а исследовать следующий слой.",
           "size": 16, "color": INK, "before": 6, "line": 1.1}])
    card(sl, 6.83, 4.7, 5.6, 1.7, fill=LIGHT)
    text(sl, 7.08, 4.95, 5.1, 1.25,
         [{"text": "Фокус", "size": 18, "bold": True, "color": NAVY},
          {"text": "Доступ к информации для человека с нарушением слуха; "
                   "и удобство для широкой аудитории.", "size": 16, "color": INK,
           "before": 6, "line": 1.1}])
    footer(sl, "Автор проекта — Никитин Даниил Константинович (AUTHOR_COUNT = 1). "
               "Подробнее: AUTHORSHIP_POLICY.md")


def s03_problem(prs):
    sl = slide(prs)
    header(sl, "Проблема", "Что происходит в большой аудитории", 30)

    box_text(sl, 0.8, 2.05, 2.6, 1.1,
             [{"text": "Преподаватель\n+ микрофон", "size": 17, "bold": True, "color": NAVY}],
             fill=WHITE, line=NAVY)
    box_text(sl, 3.9, 1.9, 5.5, 1.4,
             [{"text": "Аудитория", "size": 17, "bold": True, "color": NAVY},
              {"text": "расстояние · шум · реверберация", "size": 15, "color": INK, "before": 4}],
             fill=LIGHT, line=BLUE)
    box_text(sl, 9.9, 2.05, 2.6, 1.1,
             [{"text": "Студент:\nтеряет часть речи", "size": 17, "bold": True, "color": NAVY}],
             fill=WHITE, line=NAVY)
    arrow(sl, 3.4, 2.6, 3.9, 2.6)
    arrow(sl, 9.4, 2.6, 9.9, 2.6)

    chips = [
        ("Расстояние", "на дальних рядах речь слабее"),
        ("Шум", "разговоры, техника, зал"),
        ("Реверберация", "«эхо» съедает согласные"),
        ("Потеря части речи", "фраза пропала — вернуть нельзя"),
    ]
    for i, (t, s) in enumerate(chips):
        x = 0.62 + i * 3.1
        card(sl, x, 4.25, 2.8, 1.15)
        text(sl, x + 0.15, 4.4, 2.5, 0.9,
             [{"text": t, "size": 16, "bold": True, "color": NAVY},
              {"text": s, "size": 13, "color": MUTED, "before": 4, "line": 1.05}])

    box_text(sl, 0.62, 5.7, 11.9, 0.95,
             [{"text": "Для человека с нарушением слуха это не удобство, а доступ к информации.",
               "size": 20, "bold": True, "color": WHITE}], fill=NAVY)
    footer(sl, "Акустика класса: BB93 (шум ≤35/40 дБ); +15 дБ SNR недостаточно для младших "
               "школьников (Bradley & Sato, 2008). См. CLEAN_FEED_ASR_RESEARCH.md")


def s04_evolution(prs):
    sl = slide(prs)
    header(sl, "Предшествующие технологии", "От радиокласса к цифровой доступной аудитории", 26)

    stages = [
        ("СОНЕТ 2.0 / FM", "чистый звук в ухо\nгрупповое использование\nиндукционный выход",
         GREEN, "решает уже сегодня"),
        ("Phonak Roger", "2.4 ГГц, <20 мс\nсеть до 35 микрофонов\nNeckLoop: USB → STT",
         GREEN, "решает уже сегодня"),
        ("Auracast ALS", "открытый стандарт\nпотребительские устройства\nпарк СА/КИ пока ограничен",
         BLUE, "развивается"),
        ("Точка звука / Route+", "единая сессия\nаудио + текст\n«Не расслышал»",
         ORANGE, "исследуемый слой"),
    ]
    for i, (t, s, col, tag) in enumerate(stages):
        x = 0.7 + i * 3.08
        card(sl, x, 1.75, 2.7, 1.9, fill=WHITE)
        rect(sl, x, 1.75, 2.7, 0.16, fill=col, shape=MSO_SHAPE.RECTANGLE)
        text(sl, x + 0.14, 2.02, 2.42, 1.55,
             [{"text": t, "size": 16, "bold": True, "color": NAVY},
              {"text": s, "size": 13, "color": INK, "before": 5, "line": 1.02}])
        badge(sl, x, 3.75, 2.7, 0.35, tag, col, WHITE, 11)
        if i < 3:
            arrow(sl, x + 2.72, 2.5, x + 3.06, 2.5)

    box_text(sl, 0.7, 4.5, 11.93, 1.65,
             [{"text": "Существующие FM-, Roger- и Auracast-системы решают задачу персональной "
                       "доставки речи.", "size": 19, "bold": True, "color": WHITE},
              {"text": "В проекте исследуется следующий уровень: единая образовательная сессия, "
                       "где тот же чистый речевой источник используется для персонального аудио "
                       "и синхронизированного цифрового сопровождения пользователя.",
               "size": 16, "color": LIGHT, "before": 8, "line": 1.1}], fill=NAVY)
    footer(sl, "Источники (2026-09-26): istok-audio.com; phonak.com (Roger NeckLoop "
               "USB→STT, 2021); bluetooth.com. См. ASSISTIVE_AUDIO_EVOLUTION.md")


def s05_auracast(prs):
    sl = slide(prs)
    header(sl, "Техническая основа", "Auracast / LE Audio: кратко", 30)

    card(sl, 0.7, 1.6, 5.8, 2.75)
    text(sl, 0.95, 1.8, 5.3, 2.4,
         [{"text": "Что это", "size": 20, "bold": True, "color": NAVY},
          {"text": "• Возможность LE Audio, определяемая профилем PBP", "size": 15,
           "color": INK, "before": 8, "line": 1.1},
          {"text": "• Роли: Broadcast Source / Sink / Assistant", "size": 15, "color": INK,
           "before": 6, "line": 1.1},
          {"text": "• BASS обязателен на приёмнике (Auracast-совместимом)", "size": 15,
           "color": INK, "before": 6, "line": 1.1},
          {"text": "• Работает локально, без интернета", "size": 15, "color": INK,
           "before": 6, "line": 1.1}])

    card(sl, 6.83, 1.6, 5.8, 2.75, fill=ORANGE_L)
    text(sl, 7.08, 1.8, 5.3, 2.4,
         [{"text": "Чего это НЕ значит", "size": 20, "bold": True, "color": ORANGE},
          {"text": "• Bluetooth 5.2 ≠ Auracast: нужны PBP, стек и квалификация", "size": 15,
           "color": INK, "before": 8, "line": 1.1},
          {"text": "• «~100 м» и «неограниченно» — не гарантия", "size": 15, "color": INK,
           "before": 6, "line": 1.1},
          {"text": "• Не любые СА/КИ принимают Auracast сегодня", "size": 15, "color": INK,
           "before": 6, "line": 1.1},
          {"text": "• Телефон не входит в аудиотракт", "size": 15, "color": INK,
           "before": 6, "line": 1.1}])

    card(sl, 0.7, 4.55, 11.93, 2.15, fill=LIGHT2)
    text(sl, 0.95, 4.75, 11.4, 1.8,
         [{"text": "Параметры для планирования", "size": 18, "bold": True, "color": NAVY},
          {"text": "LC3: кадры 7.5/10 мс · базовые конфигурации 16/24 кГц (обязательны) · "
                   "48 кГц — опция", "size": 15, "color": INK, "before": 8},
          {"text": "Presentation delay: приёмник обязан поддерживать 40 мс; диапазон 20–40 мс",
           "size": 15, "color": INK, "before": 6},
          {"text": "TARGET задержки проекта: ≤60 мс; ACCEPTABLE ≤100 мс — статус TARGET, "
                   "подлежит проверке (E01)", "size": 15, "color": ORANGE, "bold": True,
           "before": 6}])
    badge(sl, 11.45, 0.78, 1.28, 0.4, "SPEC", BLUE, WHITE, 12)
    footer(sl, "Источники: Bluetooth SIG PBP 1.0.1, BAP 1.0.1, BASS 1.0; "
               "AURACAST_TECHNICAL_MAP.md; LC3_CONFIGURATION_MATRIX.md (доступ 2026-09-26)")


def s06_architecture(prs):
    sl = slide(prs)
    header(sl, "Архитектура", "«Точка звука»: один чистый источник — две ветки", 27)

    # clean source
    box_text(sl, 4.87, 1.5, 3.6, 0.6,
             [{"text": "CLEAN SOURCE (PCM)", "size": 17, "bold": True, "color": NAVY}],
             fill=WHITE, line=NAVY)
    # split
    arrow(sl, 6.67, 2.1, 6.67, 2.45)
    arrow(sl, 2.75, 2.45, 10.58, 2.45, head=False, width=1.5)
    arrow(sl, 2.75, 2.45, 2.75, 2.6)
    arrow(sl, 10.58, 2.45, 10.58, 2.6)

    box_text(sl, 1.55, 2.6, 2.4, 0.55,
             [{"text": "AUDIO BRANCH", "size": 14, "bold": True, "color": NAVY}],
             fill=LIGHT, line=BLUE)
    box_text(sl, 9.38, 2.6, 2.4, 0.55,
             [{"text": "DATA BRANCH", "size": 14, "bold": True, "color": GREEN}],
             fill=GREEN_L, line=GREEN)

    arrow(sl, 2.75, 3.15, 2.75, 3.45)
    arrow(sl, 10.58, 3.15, 10.58, 3.45)

    box_text(sl, 1.55, 3.45, 2.4, 0.6,
             [{"text": "LC3", "size": 18, "bold": True, "color": NAVY}], fill=WHITE, line=NAVY)
    box_text(sl, 9.38, 3.45, 2.4, 0.6,
             [{"text": "ASR", "size": 18, "bold": True, "color": GREEN}],
             fill=WHITE, line=GREEN)

    arrow(sl, 2.75, 4.05, 2.75, 4.35)
    arrow(sl, 10.58, 4.05, 10.58, 4.35)

    box_text(sl, 1.35, 4.35, 2.8, 0.65,
             [{"text": "Auracast эфир", "size": 18, "bold": True, "color": NAVY}],
             fill=LIGHT, line=BLUE)
    box_text(sl, 9.38, 4.35, 2.4, 0.65,
             [{"text": "Transcript", "size": 18, "bold": True, "color": GREEN}],
             fill=GREEN_L, line=GREEN)

    # audio split to Direct / Bridge-T
    arrow(sl, 2.75, 5.0, 2.75, 5.2, head=False)
    arrow(sl, 1.75, 5.2, 4.15, 5.2, head=False)
    arrow(sl, 1.75, 5.2, 1.75, 5.35)
    arrow(sl, 4.15, 5.2, 4.15, 5.35)
    box_text(sl, 0.75, 5.35, 2.0, 0.6,
             [{"text": "Direct sink", "size": 15, "bold": True, "color": NAVY}],
             fill=WHITE, line=NAVY)
    box_text(sl, 3.15, 5.35, 2.0, 0.6,
             [{"text": "Bridge-T", "size": 15, "bold": True, "color": NAVY}],
             fill=WHITE, line=NAVY)
    arrow(sl, 1.75, 5.95, 1.75, 6.2)
    arrow(sl, 4.15, 5.95, 4.15, 6.2)
    box_text(sl, 0.75, 6.2, 2.0, 0.55,
             [{"text": "Наушники / СА / КИ", "size": 12, "color": INK}],
             fill=LIGHT2, line=LINEC)
    box_text(sl, 3.15, 6.2, 2.0, 0.55,
             [{"text": "T/MT → СА / КИ", "size": 12, "color": INK}],
             fill=LIGHT2, line=LINEC)

    # data to app
    arrow(sl, 10.58, 5.0, 10.58, 5.35)
    box_text(sl, 9.38, 5.35, 2.4, 0.6,
             [{"text": "Route+ App", "size": 16, "bold": True, "color": ORANGE}],
             fill=ORANGE_L, line=ORANGE)
    text(sl, 8.3, 6.1, 4.5, 0.7,
         [{"text": "Route+ — текстовый слой; не часть аудиоцепочки Auracast",
           "size": 12, "italic": True, "color": MUTED, "line": 1.05}])

    # legend
    rect(sl, 0.7, 1.62, 0.14, 0.14, fill=BLUE, shape=MSO_SHAPE.RECTANGLE)
    text(sl, 0.9, 1.56, 1.2, 0.25, [{"text": "AUDIO", "size": 11, "color": INK}])
    rect(sl, 2.2, 1.62, 0.14, 0.14, fill=GREEN, shape=MSO_SHAPE.RECTANGLE)
    text(sl, 2.4, 1.56, 1.2, 0.25, [{"text": "DATA", "size": 11, "color": INK}])
    rect(sl, 3.7, 1.62, 0.14, 0.14, fill=ORANGE, shape=MSO_SHAPE.RECTANGLE)
    text(sl, 3.9, 1.56, 1.6, 0.25, [{"text": "TEXT (APP)", "size": 11, "color": INK}])
    footer(sl, "Схема концептуальная; числовые характеристики намеренно не указаны. "
               "См. SYSTEM_ARCHITECTURE.md (GATE 6A)")


def s07_audio_paths(prs):
    sl = slide(prs)
    header(sl, "Доставка звука", "Кому и как передаётся звук", 30)

    cols = [
        ("Прямой Auracast", "парк ограничен", GREEN,
         ["Auracast эфир", "Прямой приём", "Наушники / совместимые СА / КИ"]),
        ("Bridge-T", "известный класс", ORANGE,
         ["Auracast sink", "LC3 → PCM → DAC → усилитель → катушка", "T/MT → СА / КИ"]),
        ("Только текст", "MVP", BLUE,
         ["Clean feed (PCM)", "ASR", "Route+ App: текст, заметки, доступность"]),
    ]
    for i, (title, tag, col, steps) in enumerate(cols):
        x = 0.7 + i * 4.15
        card(sl, x, 1.6, 3.8, 4.55)
        text(sl, x + 0.2, 1.74, 3.4, 0.4,
             [{"text": title, "size": 19, "bold": True, "color": NAVY}])
        badge(sl, x + 0.2, 2.2, 2.1, 0.34, tag, col, WHITE, 11)
        y = 2.78
        for j, st in enumerate(steps):
            box_text(sl, x + 0.35, y, 3.1, 0.8,
                     [{"text": st, "size": 14, "bold": j == 0, "color": INK, "line": 1.05}],
                     fill=LIGHT2, line=col if j < 2 else LINEC)
            if j < len(steps) - 1:
                arrow(sl, x + 1.9, y + 0.8, x + 1.9, y + 1.05, color=col)
            y += 1.05
    footer(sl, "Bridge-T — известный класс (Roger MyLink/NeckLoop, Auri RX1, Bettear RTX, "
               "AuraCoil); конкретные технические отличия — предмет исследования. "
               "См. AURACAST_TELECOIL_PRIOR_ART.md")


def s08_app(prs):
    sl = slide(prs)
    header(sl, "Route+ App", "Ассистивный текстовый слой (реальный MVP)", 30)
    badge(sl, 6.1, 1.5, 2.2, 0.4, "IMPLEMENTED", GREEN, WHITE, 12)

    if os.path.exists(os.path.join(SHOTS, "01_session.png")):
        picture_cover(sl, os.path.join(SHOTS, "01_session.png"), 0.7, 1.5, 2.35, 5.0)
    text(sl, 3.25, 1.5, 2.9, 0.4, [{"text": "Экран лекции", "size": 12, "color": MUTED}])

    text(sl, 6.1, 2.05, 6.6, 2.6,
         [{"text": "• Вход в сессию (QR) и статус трансляции", "size": 16, "color": INK,
           "after": 5, "line": 1.05},
          {"text": "• Live captions: partial и final различаются", "size": 16, "color": INK,
           "after": 5, "line": 1.05},
          {"text": "• Таймкоды и история текста лекции", "size": 16, "color": INK,
           "after": 5, "line": 1.05},
          {"text": "• «Не расслышал», заметки, закладки", "size": 16, "color": INK,
           "after": 5, "line": 1.05},
          {"text": "• Доступность: размер, контраст, анимация, вибрация", "size": 16,
           "color": INK, "line": 1.05}])

    if os.path.exists(os.path.join(SHOTS, "02_live_captions.png")):
        picture_cover(sl, os.path.join(SHOTS, "02_live_captions.png"), 6.1, 4.55, 6.6, 2.0)
        text(sl, 6.1, 6.6, 6.6, 0.3,
             [{"text": "Фрагмент экрана с живыми субтитрами", "size": 11, "color": MUTED}])
    footer(sl, "Скриншоты получены из работающего MVP (React + FastAPI, demo mode). "
               "Аудио в MVP не воспроизводится. См. 06_application/README.md")


def s09_missed(prs):
    sl = slide(prs)
    header(sl, "Ключевая функция", "«Не расслышал»: вернуть пропущенное", 30)

    steps = ["Преподаватель говорит", "Студент пропустил фразу",
             "Нажал «Не расслышал»", "Получил последние 15 секунд текста"]
    for i, st in enumerate(steps):
        x = 0.7 + i * 3.08
        box_text(sl, x, 1.6, 2.7, 1.0,
                 [{"text": st, "size": 14, "bold": True, "color": NAVY, "line": 1.05}],
                 fill=LIGHT if i < 3 else ORANGE_L, line=BLUE if i < 3 else ORANGE)
        if i < 3:
            arrow(sl, x + 2.72, 2.1, x + 3.06, 2.1, head=True)

    if os.path.exists(os.path.join(SHOTS, "03_missed_speech.png")):
        picture_cover(sl, os.path.join(SHOTS, "03_missed_speech.png"), 0.85, 2.9, 2.15, 3.9)

    text(sl, 3.35, 2.9, 9.3, 3.4,
         [{"text": "• Дословный распознанный текст — без AI-пересказа", "size": 17,
           "color": INK, "after": 8, "line": 1.05},
          {"text": "• Из окна можно сохранить фрагмент или добавить заметку", "size": 17,
           "color": INK, "after": 8, "line": 1.05},
          {"text": "• Частичный (partial) текст визуально отличается от финального",
           "size": 17, "color": INK, "after": 8, "line": 1.05},
          {"text": "• 15 секунд — MVP default, а не доказанный оптимум: окно исследуется "
                   "в E06/E10", "size": 17, "color": ORANGE, "bold": True, "line": 1.05}])
    footer(sl, "Ближайший предшественник механики: Bionica US20090076804 (2007, buffer + "
               "instant replay + STT). См. MISSED_SPEECH_PRIOR_ART.md")


def s10_analogs(prs):
    sl = slide(prs)
    header(sl, "Аналоги", "Сравнение по значимым критериям", 30)

    systems = ["СОНЕТ 2.0", "Roger +\nNeckLoop", "Auri /\nBettear", "HearAura",
               "Точка звука\n[PROPOSED]"]
    rows = [
        ("Чистый звук в ухо", ["YES", "YES", "YES", "NO", "PROPOSED"]),
        ("Текст из того же источника", ["NO", "PARTIAL\n(USB → STT)", "PARTIAL", "YES", "YES"]),
        ("«Не расслышал»", ["NO", "NO", "NO", "PARTIAL\n(summary)", "YES"]),
        ("Потребительские устройства", ["NO", "NO", "YES", "YES", "YES"]),
        ("Заметки и доступность", ["NO", "PARTIAL", "NO", "YES", "YES"]),
    ]
    x0, y0 = 0.3, 1.6
    fw, cw, rh = 2.93, 1.96, 0.62
    # header
    box_text(sl, x0, y0, fw, 0.7, [{"text": "Критерий", "size": 13, "bold": True,
                                    "color": WHITE}], fill=NAVY)
    for j, s in enumerate(systems):
        box_text(sl, x0 + fw + j * cw, y0, cw, 0.7,
                 [{"text": s, "size": 12, "bold": True, "color": WHITE, "line": 1.0}],
                 fill=NAVY)
    for i, (feat, vals) in enumerate(rows):
        y = y0 + 0.7 + i * rh
        box_text(sl, x0, y, fw, rh, [{"text": feat, "size": 13, "bold": True, "color": NAVY}],
                 fill=LIGHT2, line=LINEC, align=PP_ALIGN.LEFT)
        for j, v in enumerate(vals):
            color = GREEN if v.startswith("YES") else (ORANGE if v.startswith("PARTIAL")
                                                       else (BLUE if v.startswith("PROPOSED")
                                                             else MUTED))
            box_text(sl, x0 + fw + j * cw, y, cw, rh,
                     [{"text": v, "size": 12, "bold": v.startswith("PROPOSED"), "color": color,
                       "line": 1.0}], fill=WHITE, line=LINEC)

    text(sl, 0.3, 5.52, 12.7, 1.0,
         [{"text": "Легенда: YES — функция есть · NO — нет · PARTIAL — частично/через внешний "
                   "компонент · PROPOSED — предлагается в проекте и ещё не реализовано.",
           "size": 13, "color": MUTED, "line": 1.1},
          {"text": "Источники (2026-09-26): istok-audio.com; phonak.com (Roger NeckLoop); "
                   "ampetronic.com; bettear.com; hearaura.app; ava.me; Google Live Transcribe docs.",
           "size": 11, "color": MUTED, "before": 6, "line": 1.1}])
    footer(sl, "Полная матрица — COMPETITOR_MATRIX.md. Таблица не подгоняется в пользу проекта.")


def s11_difference(prs):
    sl = slide(prs)
    header(sl, "Отличие проекта", "Что именно исследуется", 30)

    box_text(sl, 1.2, 1.6, 10.93, 1.35,
             [{"text": "Отдельные элементы известны. Исследуемый слой — единая "
                       "образовательная сессия.", "size": 24, "bold": True, "color": WHITE}],
             fill=NAVY)

    chips = [
        ("Общий временной контекст", "аудио и текст в одном таймлайне сессии"),
        ("Audio + synchronized data", "тот же источник: звук и текстовые события"),
        ("Missed Speech UX", "дословный фрагмент последних N секунд"),
        ("Confidence-aware", "мягкая индикация неуверенного распознавания (research)"),
        ("Единый accessibility layer", "direct Auracast + Bridge-T + text-only"),
        ("Привязка к сессии", "QR → сессия; текст связан с конкретной лекцией"),
    ]
    for i, (t, s) in enumerate(chips):
        col = i % 3
        row = i // 3
        x = 0.7 + col * 4.15
        y = 3.2 + row * 1.3
        card(sl, x, y, 3.8, 1.1)
        text(sl, x + 0.18, y + 0.12, 3.45, 0.9,
             [{"text": t, "size": 15, "bold": True, "color": NAVY},
              {"text": s, "size": 12, "color": MUTED, "before": 3, "line": 1.0}])

    box_text(sl, 0.7, 5.95, 11.93, 0.8,
             [{"text": "Clean-feed → STT известен (Roger NeckLoop USB→STT, 2021). "
                       "Патентная новизна не заявляется до профессионального поиска (GATE 7).",
               "size": 15, "bold": True, "color": ORANGE}], fill=ORANGE_L, line=ORANGE)
    footer(sl, "См. NOVELTY_MAP.md (N02 NOT_NEW; N05 KNOWN_CLASS; N09 INSUFFICIENT_EVIDENCE) "
               "и DIFFERENTIATION_STATEMENT.md")


def s12_experiments(prs):
    sl = slide(prs)
    header(sl, "Научно-инженерная часть", "Эксперименты: что и как будем измерять", 28)

    exps = [
        ("E01", "Latency", "сквозная задержка «микрофон → выход»",
         "TARGET ≤60 мс; ≤100 допустимо; >150 — брак"),
        ("E03", "Stability", "устойчивость потока ≥1 ч",
         "нет разрывов в сценарии лекции"),
        ("E04", "LC3 configs", "сравнение 16_2_1 / 24_2_1 / HQ",
         "выбор baseline по качеству и задержке"),
        ("E05", "Clean-feed vs phone ASR", "WER/CER, условия A–E",
         "измеримое преимущество или его отсутствие"),
        ("E07", "Bridge-T", "Auracast → индукция → T/MT",
         "разборчивый сигнал без гула"),
    ]
    for i, (code, title, what, acc) in enumerate(exps):
        x = 0.4 + i * 2.57
        card(sl, x, 1.6, 2.3, 3.9)
        badge(sl, x + 0.35, 1.78, 0.85, 0.38, code, NAVY, WHITE, 13)
        badge(sl, x + 1.3, 1.78, 0.85, 0.38, "PLANNED", ORANGE, WHITE, 10)
        text(sl, x + 0.18, 2.35, 1.95, 3.0,
             [{"text": title, "size": 16, "bold": True, "color": NAVY, "line": 1.0},
              {"text": what, "size": 13, "color": INK, "before": 6, "line": 1.05},
              {"text": "Критерий:", "size": 11, "bold": True, "color": MUTED, "before": 10},
              {"text": acc, "size": 13, "color": ORANGE, "line": 1.05}])

    box_text(sl, 0.4, 5.7, 12.53, 0.8,
             [{"text": "EXPERIMENT PLANNED ≠ RESULT ACHIEVED. Ни одна характеристика "
                       "не измерена; все значения — целевые.", "size": 17, "bold": True,
               "color": WHITE}], fill=ORANGE)
    footer(sl, "Протоколы: EXPERIMENT_READINESS_GATE6.md; LATENCY_TEST_PROTOCOL.md; "
               "ASR_COMPARISON_PROTOCOL.md")


def s13_hardware(prs):
    sl = slide(prs)
    header(sl, "Hardware MVP", "Стенд для экспериментов", 30)
    badge(sl, 9.0, 1.52, 4.0, 0.42, "RECOMMENDED RESEARCH STAND", ORANGE, WHITE, 12)
    badge(sl, 9.0, 2.02, 4.0, 0.42, "НЕ ЗАКУПЛЕНО", WHITE, ORANGE, 12)
    rect(sl, 9.0, 2.02, 4.0, 0.42, fill=None, line=ORANGE, lw=1.5, radius=0.5)

    boxes = [
        "Источник\n(ПК / микрофон)",
        "nRF5340 Audio DK\nBroadcast SOURCE",
        "Auracast\nэфир",
        "nRF5340 Audio DK\nBroadcast SINK",
        "Аналоговый выход\n(3.5 мм)",
    ]
    xs = [0.55, 3.05, 5.55, 8.05, 10.55]
    for i, (x, b) in enumerate(zip(xs, boxes)):
        col = BLUE if i in (1, 3) else NAVY
        box_text(sl, x, 2.6, 2.2, 1.0,
                 [{"text": b, "size": 13, "bold": True, "color": col, "line": 1.05}],
                 fill=LIGHT if i in (1, 3) else WHITE, line=col)
        if i < 4:
            arrow(sl, x + 2.2, 3.1, x + 2.5, 3.1, color=NAVY)

    arrow(sl, 9.15, 3.6, 9.15, 3.95, color=ORANGE)
    box_text(sl, 7.4, 3.95, 3.5, 0.7,
             [{"text": "Bridge-T branch: индукция → T/MT", "size": 14, "bold": True,
               "color": ORANGE}], fill=ORANGE_L, line=ORANGE)
    arrow(sl, 9.15, 4.65, 9.15, 5.0, color=ORANGE)
    box_text(sl, 7.4, 5.0, 3.5, 0.6,
             [{"text": "СА / КИ (старый парк)", "size": 13, "color": INK}],
             fill=LIGHT2, line=LINEC)

    text(sl, 0.55, 5.0, 6.5, 1.3,
         [{"text": "Что даёт стенд", "size": 17, "bold": True, "color": NAVY},
          {"text": "• Контроль LC3-конфигураций и QoS", "size": 15, "color": INK,
           "before": 6, "line": 1.05},
          {"text": "• Управляемые измерения задержки и устойчивости", "size": 15,
           "color": INK, "before": 4, "line": 1.05},
          {"text": "• Low-cost demo: FMA120 + AuraClip (<$150–250)", "size": 13,
           "color": MUTED, "before": 6, "line": 1.05}])
    footer(sl, "Альтернативы и обоснование — PURCHASE_DECISION.md; DEVKIT_COMPARISON.md. "
               "Стенд на devkit ≠ квалифицированный продукт.")


def s14_deployments(prs):
    sl = slide(prs)
    header(sl, "Реальные внедрения", "Auracast ALS в образовании", 30)

    cards = [
        ("University of Queensland", "PRIMARY", GREEN,
         "65 лекционных аудиторий.\nТрансмиттеры Audeara, выдача приёмников, QR-инструкции.",
         "Официально: «iPhones do not currently support this feature»."),
        ("University of Oxford", "vendor + press", ORANGE,
         "Библиотека Bodleian → стандарт кампуса.\nAuri TX2N-D (Dante), 16× RX1, neckloops.",
         "Подтверждено вендором и отраслевой прессой; ox.ac.uk недоступен."),
        ("UAL, Creative Computing Institute", "PRIMARY", GREEN,
         "3 аудитории (2×96, 1×48).\nAuri TX2N-D, 4× RX1; монтаж ~1 час.",
         "Малый пилот; для устройств без Auracast — выдача приёмников."),
    ]
    for i, (org, conf, col, body, note) in enumerate(cards):
        x = 0.7 + i * 4.11
        card(sl, x, 1.7, 3.75, 3.9)
        text(sl, x + 0.2, 1.9, 3.35, 0.8,
             [{"text": org, "size": 16, "bold": True, "color": NAVY, "line": 1.0}])
        badge(sl, x + 0.2, 2.75, 1.75, 0.36, conf, col, WHITE, 10)
        text(sl, x + 0.2, 3.25, 3.35, 1.4,
             [{"text": body, "size": 14, "color": INK, "line": 1.1}])
        text(sl, x + 0.2, 4.75, 3.35, 0.75,
             [{"text": note, "size": 11, "italic": True, "color": MUTED, "line": 1.05}])

    box_text(sl, 0.7, 5.85, 11.93, 0.75,
             [{"text": "Развёртывания подтверждают аудиотракт. Live captions из того же "
                       "источника в реальных внедрениях не подтверждены.", "size": 15,
               "bold": True, "color": NAVY}], fill=LIGHT)
    footer(sl, "Источники (2026-09-26): news.uq.edu.au; my.uq.edu.au; auriaudio.com; "
               "listentech.com; bluetooth.com location profiles; AV Magazine. "
               "См. REAL_WORLD_DEPLOYMENTS.md")


def s15_value(prs):
    sl = slide(prs)
    header(sl, "Ценность", "Для кого и что меняется", 30)

    box_text(sl, 0.7, 1.75, 5.8, 3.5,
             [{"text": "Для человека\nс нарушением слуха", "size": 24, "bold": True,
               "color": WHITE},
              {"text": "Доступ к информации —\nне удобство, а необходимость.", "size": 17,
               "color": LIGHT, "before": 14, "line": 1.15},
              {"text": "• персональный звук\n• текст из того же источника\n"
                       "• «Не расслышал»\n• поддержка старых СА/КИ", "size": 15,
               "color": WHITE, "before": 14, "line": 1.2}], fill=NAVY)

    box_text(sl, 6.83, 1.75, 5.8, 3.5,
             [{"text": "Для широкой\nаудитории", "size": 24, "bold": True, "color": NAVY},
              {"text": "Разборчивость и текст там,\nгде акустика не помогает.", "size": 17,
               "color": INK, "before": 14, "line": 1.15},
              {"text": "• дальние ряды\n• шумные и гулкие залы\n"
                       "• иностранные студенты\n• заметки и конспект", "size": 15,
               "color": INK, "before": 14, "line": 1.2}], fill=LIGHT2, line=LINEC)

    box_text(sl, 0.7, 5.55, 11.93, 0.95,
             [{"text": "Для широкой аудитории это удобство. Для человека с нарушением "
                       "слуха — доступ к информации.", "size": 20, "bold": True,
               "color": NAVY}], fill=LIGHT, line=BLUE)
    footer(sl, "Принцип проекта. Формулировка «здоровые люди» не используется.")


def s16_roadmap(prs):
    sl = slide(prs)
    header(sl, "Roadmap", "Что сделано, что дальше", 30)

    cols = [
        ("01 · NOW", GREEN, "Route+ Software MVP",
         ["• Live captions, «Не расслышал», заметки, доступность",
          "• Demo mode без железа",
          "• Тесты: 19 backend + 4 frontend"]),
        ("02 · NEXT", ORANGE, "Hardware stand + эксперименты",
         ["• 2× nRF5340 Audio DK (закупка — по решению)",
          "• E01 latency · E03 stability · E04 LC3",
          "• E05 clean-feed vs phone ASR · E07 Bridge-T"]),
        ("03 · THEN", BLUE, "Пилот и расширение",
         ["• Пилот в учебной аудитории (E08/E09/E10)",
          "• Уточнение окна «Не расслышал»",
          "• AI-слой: перевод/конспект — FUTURE"]),
    ]
    for i, (tag, col, title, bullets) in enumerate(cols):
        x = 0.7 + i * 4.11
        card(sl, x, 1.75, 3.75, 3.8, fill=WHITE)
        rect(sl, x, 1.75, 3.75, 0.2, fill=col, shape=MSO_SHAPE.RECTANGLE)
        text(sl, x + 0.22, 2.1, 3.3, 0.4,
             [{"text": tag, "size": 14, "bold": True, "color": col}])
        text(sl, x + 0.22, 2.5, 3.3, 0.7,
             [{"text": title, "size": 17, "bold": True, "color": NAVY, "line": 1.0}])
        text(sl, x + 0.22, 3.3, 3.35, 2.1,
             [{"text": b, "size": 14, "color": INK, "after": 6, "line": 1.08} for b in bullets])
        if i < 2:
            arrow(sl, x + 3.78, 3.6, x + 4.08, 3.6, color=col)

    footer(sl, "Статусы: NOW — реализовано (MVP); NEXT — запланировано, не выполнено; "
               "THEN — после экспериментов. См. ROADMAP.md")


def s17_closing(prs):
    sl = slide(prs)
    rect(sl, 0, 0, SW, SH, fill=NAVY, shape=MSO_SHAPE.RECTANGLE)
    text(sl, M, 0.7, SW - 2 * M, 0.8,
         [{"text": "Итог", "size": 34, "bold": True, "color": WHITE}])
    text(sl, M, 1.75, 11.9, 2.6,
         [{"text": "• Реализован программный MVP ассистивного текстового слоя.",
           "size": 18, "color": LIGHT, "after": 10, "line": 1.1},
          {"text": "• Определён архитектурный baseline: чистый источник → аудио (Auracast / "
                   "Bridge-T) и текст (ASR → Route+).", "size": 18, "color": LIGHT,
           "after": 10, "line": 1.1},
          {"text": "• Подготовлены протоколы экспериментов; hardware-стенд рекомендован.",
           "size": 18, "color": LIGHT, "after": 10, "line": 1.1},
          {"text": "• Новизна не заявляется: известные блоки названы честно; отличие — "
                   "в исследуемом системном слое.", "size": 18, "color": LIGHT,
           "line": 1.1}])

    badge(sl, M, 4.6, 3.5, 0.45, "IMPLEMENTED — MVP", GREEN, WHITE, 13)
    badge(sl, 4.3, 4.6, 3.8, 0.45, "PLANNED — эксперименты", ORANGE, WHITE, 13)
    badge(sl, 8.3, 4.6, 3.5, 0.45, "FUTURE — AI-слой", BLUE, WHITE, 13)

    text(sl, M, 5.6, 11.9, 0.9,
         [{"text": "Автор: Никитин Даниил Константинович (AUTHOR_COUNT = 1)", "size": 16,
           "bold": True, "color": WHITE},
          {"text": "Материалы и источники: PROJECT_STATE.md, MASTER_REPORT.md, "
                   "research-файлы GATE 2–6A.", "size": 12, "color": BLUE_SOFT,
           "before": 6, "line": 1.1}])


def main():
    prs = new_deck()
    for fn in (s01_title, s02_motivation, s03_problem, s04_evolution, s05_auracast,
               s06_architecture, s07_audio_paths, s08_app, s09_missed, s10_analogs,
               s11_difference, s12_experiments, s13_hardware, s14_deployments,
               s15_value, s16_roadmap, s17_closing):
        fn(prs)
    prs.save(OUT)
    print("saved:", OUT)
    print("slides:", len(prs.slides))


if __name__ == "__main__":
    main()
